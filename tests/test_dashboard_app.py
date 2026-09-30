"""Exercise the actual Streamlit controls and generated outputs."""
from pathlib import Path
import unittest
from unittest.mock import patch

from streamlit.testing.v1 import AppTest
from tests.streamlit_helpers import stateful_tabs


class DashboardAppTests(unittest.TestCase):
    def setUp(self):
        transport = stateful_tabs(); transport.start(); self.addCleanup(transport.stop)
        self.app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py"), default_timeout=15).run()
        self.assertEqual(len(self.app.exception), 0)

    def page(self, name):
        self.app.session_state["report_navigation"] = name
        self.app.run()
        self.assertEqual(len(self.app.exception), 0)

    def test_initial_dashboard_and_chat(self):
        self.assertEqual([tab.label for tab in self.app.tabs], ["Reporting", "Overview", "Ledger explorer", "Review & export", "Executive brief", "Planning", "Budgeting", "Forecasting", "Business drivers", "Cash planning", "Plan review", "Guide"])
        self.app.button(key="chat_suggestion_0").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.chat_message), 2)
        self.assertIn("Local data analysis", [c.value for c in self.app.caption if "Local data analysis" in c.value][0])

    def test_filter_changes_clear_brief_and_chat(self):
        self.app.button(key="chat_suggestion_0").click().run()
        self.page("Executive brief")
        [b for b in self.app.button if b.label == "Generate executive brief"][0].click().run()
        self.assertTrue(self.app.session_state["executive_markdown"])
        self.app.multiselect(key="dashboard_cost_center").set_value(["Finance"]).run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.session_state["chat_messages"], [])
        self.assertNotIn("executive_markdown", self.app.session_state)
        self.page("Ledger explorer")
        self.assertEqual(len(self.app.dataframe[0].value), 120)
        self.app.button(key="chat_suggestion_0").click().run()
        self.assertIn("Current dashboard view", self.app.session_state["chat_messages"][-1]["scope"])

    def test_empty_intersection_disables_data_actions(self):
        self.app.multiselect(key="dashboard_period").set_value([12])
        self.app.multiselect(key="dashboard_quarter").set_value(["Q1"]).run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(self.app.button(key="chat_suggestion_0").disabled)
        self.assertTrue(self.app.chat_input[0].disabled)
        self.page("Executive brief")
        self.assertTrue([b for b in self.app.button if b.label == "Generate executive brief"][0].disabled)
        self.app.button(key="reset_dashboard_filters").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertFalse(self.app.chat_input[0].disabled)
        self.page("Ledger explorer")
        self.assertEqual(len(self.app.dataframe[0].value), 600)

    def test_chat_input_and_excel_download(self):
        self.app.multiselect(key="dashboard_cost_center").set_value(["Finance"]).run()
        self.app.chat_input[0].set_value("how much did we spend?").run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.chat_message), 2)
        self.app.button(key="chat_suggestion_3").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.session_state["chat_exports"]), 1)
        self.app.selectbox(key="chat_scope").select("Full ledger").run()
        self.assertEqual(self.app.session_state["chat_messages"], [])
        self.assertEqual(self.app.session_state["chat_exports"], {})

    def test_upload_mode_does_not_show_previous_demo_results(self):
        self.app.radio(key="data_mode").set_value("Upload your data").run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.chat_input), 0)
        self.assertEqual(len(self.app.dataframe), 0)

    def test_csv_and_xlsx_uploads_and_demo_switch(self):
        from io import BytesIO
        from data_engine import generate_synthetic_budget_actuals
        frame = generate_synthetic_budget_actuals().iloc[:3]
        self.app.radio(key="data_mode").set_value("Upload your data")
        for extension in ["csv", "xlsx"]:
            upload = BytesIO()
            if extension == "csv":
                upload.write(frame.to_csv(index=False).encode())
            else:
                frame.to_excel(upload, index=False)
            upload.name = "uploaded." + extension
            self.app.session_state["report_navigation"] = "Ledger explorer"
            with patch("streamlit.sidebar.file_uploader", return_value=upload):
                self.app.run()
            self.assertEqual(len(self.app.exception), 0)
            self.assertEqual(len(self.app.dataframe[0].value), 3)
            self.assertEqual(self.app.session_state["data_source"], upload.name)
        self.app.radio(key="data_mode").set_value("Demo dataset").run()
        self.assertEqual(len(self.app.exception), 0)
        self.page("Ledger explorer")
        self.assertEqual(len(self.app.dataframe[0].value), 600)

    def test_invalid_upload_does_not_display_stale_data(self):
        from io import BytesIO
        upload = BytesIO(b"wrong,columns\n1,2\n")
        upload.name = "invalid.csv"
        self.app.radio(key="data_mode").set_value("Upload your data")
        with patch("streamlit.sidebar.file_uploader", return_value=upload):
            self.app.run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 1)
        self.assertEqual(len(self.app.dataframe), 0)

    def test_reporting_cutoff_limits_every_view(self):
        self.app.selectbox(key="closed_period_2026").select(3).run()
        self.assertEqual(len(self.app.exception),0)
        self.page("Ledger explorer")
        self.assertEqual(len(self.app.dataframe[0].value),150)
        self.assertEqual(set(self.app.session_state['raw_df'].period),{1,2,3})

    def test_review_editor_saves_notes_and_builds_report_pack(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.page('Review & export')
        key=next(k for k in self.app.session_state.filtered_state if k.startswith('review_editor_'))
        self.app.session_state[key]={'edited_rows':{0:{'owner':'FP&A lead','status':'Resolved','explanation':'Matched invoice to approved budget change'}},'added_rows':[],'deleted_rows':[]}
        next(b for b in self.app.button if b.label=='Save review notes').click().run()
        self.assertEqual(len(self.app.exception),0)
        notes=self.app.session_state['review_notes']
        self.assertTrue(any(n['owner']=='FP&A lead' and n['status']=='Resolved' for n in notes.values()))
        next(b for b in self.app.button if b.label=='Prepare reporting pack').click().run()
        self.assertEqual(len(self.app.exception),0)
        book=load_workbook(BytesIO(self.app.session_state['review_pack']),data_only=True)
        self.assertIn('Review Register',book.sheetnames)
        self.app.multiselect(key='dashboard_cost_center').set_value(['Finance']).run()
        self.assertEqual(self.app.session_state['review_notes'],notes)
        self.assertNotIn('review_pack',self.app.session_state)

    def test_chat_session_bounds_memory_and_removes_expired_exports(self):
        self.app.session_state['chat_messages']=[{'role':'user' if i%2==0 else 'assistant','content':'Previous message','export_id':f'e{i}'} for i in range(40)]
        self.app.session_state['chat_exports']={f'e{i}':{'bytes':b'example','filename':'sample.xlsx','mime':'application/octet-stream'} for i in range(40)}
        self.app.button(key='chat_suggestion_0').click().run()
        self.assertEqual(len(self.app.exception),0)
        self.assertEqual(len(self.app.session_state['chat_messages']),40)
        self.assertEqual(len(self.app.session_state['chat_exports']),5)
        self.assertNotIn('e0',self.app.session_state['chat_exports'])

    def test_chat_expansion_preserves_scope_history_and_downloads(self):
        self.app.multiselect(key='dashboard_cost_center').set_value(['Finance']).run()
        self.app.selectbox(key='chat_scope').select('Full ledger').run()
        self.app.button(key='chat_suggestion_3').click().run()
        history=list(self.app.session_state['chat_messages'])
        exports=dict(self.app.session_state['chat_exports'])
        self.app.button(key='toggle_copilot_layout').click().run()
        self.assertEqual(len(self.app.exception),0)
        self.assertTrue(self.app.session_state['copilot_expanded'])
        self.assertEqual(self.app.selectbox(key='chat_scope').value,'Full ledger')
        self.assertEqual(self.app.session_state['chat_messages'],history)
        self.assertEqual(self.app.session_state['chat_exports'],exports)
        self.app.chat_input[0].set_value('total variance').run()
        self.assertEqual(len(self.app.session_state['chat_messages']),4)
        self.app.button(key='toggle_copilot_layout').click().run()
        self.assertEqual(len(self.app.exception),0)
        self.assertFalse(self.app.session_state['copilot_expanded'])
        self.assertEqual(self.app.multiselect(key='dashboard_cost_center').value,['Finance'])
        self.assertEqual(len(self.app.session_state['chat_messages']),4)
        self.assertEqual(self.app.session_state['chat_exports'],exports)
