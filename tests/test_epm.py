"""Financial planning invariants, source boundaries, and durable revision controls."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

from epm import (cents, base_plan, demo_plan, validate_plan, frame, from_frame, read_planning_file,
    from_expense_ledger, monthly_report, cash_report, variance_report, forecast, allocate_budget,
    adjust, apply_drivers, apply_actuals, replace_grid, compare_plans, backtest, roll_horizon,
    planning_answer, planning_pack, fingerprint, revenue_driver)
from epm_store import PlanStore, ConflictError
from data_engine import generate_synthetic_budget_actuals


def small_plan():
    plan = base_plan("Test plan", count=6, cutoff="2026-03")
    for number in range(1, 7):
        period = f"2026-{number:02d}"
        for category, amount in [("Revenue", 1000), ("COGS", 400), ("Payroll", 100), ("Depreciation", 20)]:
            plan["rows"].append(dict(period=period, department="Operations", account=category, category=category,
                budget_cents=cents(amount), actual_cents=cents(amount) if number <= 3 else None, forecast_cents=cents(amount)))
    plan["cash"].update(opening_cash_cents=cents(1000), opening_ar_cents=cents(200), opening_ap_cents=cents(100))
    return validate_plan(plan)


class PlanningCalculations(unittest.TestCase):
    def test_rounding_and_nonfinite_input(self):
        self.assertEqual(cents("1.005"), 101)
        self.assertEqual(cents("-1.005"), -101)
        for value in [None, True, "nan", "inf", "-inf", "1e1000"]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                cents(value)

    def test_invalid_plan_structure(self):
        for mutation in [lambda p: p["rows"].append(p["rows"][0]),
                         lambda p: p["rows"][0].update(budget_cents=True),
                         lambda p: p["rows"][0].update(actual_cents=float("nan")),
                         lambda p: p["rows"][-1].update(actual_cents=0),
                         lambda p: p["rows"].pop(),
                         lambda p: p["cash"].update(collection_lag=-1)]:
            plan = small_plan(); mutation(plan)
            with self.assertRaises(ValueError):
                validate_plan(plan)

    def test_actual_forecast_splice_and_sign_semantics(self):
        plan = small_plan()
        plan["rows"][0]["actual_cents"] = cents(1100)
        plan["rows"][1]["actual_cents"] = cents(450)
        result = monthly_report(plan)
        self.assertEqual(result.iloc[0].outlook_operating_profit, 530)
        self.assertEqual(result.iloc[0].profit_vs_budget, 50)
        self.assertEqual(result.iloc[3].outlook_operating_profit, 480)
        detail = variance_report(plan)
        self.assertEqual(detail.iloc[0].favorable_impact, 100)
        self.assertEqual(detail.iloc[1].favorable_impact, -50)

    def test_missing_actuals_do_not_become_zero(self):
        plan = small_plan(); plan["rows"][1]["actual_cents"] = None
        monthly = monthly_report(plan)
        self.assertTrue(pd.isna(monthly.iloc[0].outlook_operating_profit))
        self.assertTrue(pd.isna(monthly.outlook_operating_profit.sum(skipna=False)))

    def test_grid_preserves_actuals_and_other_inputs(self):
        plan = small_plan(); edited = frame(plan)
        edited.loc[edited.period > "2026-03", "forecast"] = 777
        result = replace_grid(plan, edited)
        self.assertEqual(result["rows"][-1]["forecast_cents"], 77700)
        self.assertEqual([r["actual_cents"] for r in result["rows"]], [r["actual_cents"] for r in plan["rows"]])
        edited.loc[0, "actual"] = 123
        with self.assertRaisesRegex(ValueError, "protected"):
            replace_grid(plan, edited)

    def test_forecast_uses_only_closed_history(self):
        plan = small_plan()
        for row in plan["rows"]:
            if row["period"] > plan["closed_through"]:
                row["forecast_cents"] = 99999999
        before = deepcopy(plan)
        result = forecast(plan, "Trailing 3 months", 10)
        self.assertEqual(result["rows"][-4]["forecast_cents"], 110000)
        self.assertEqual(result["rows"][:12], plan["rows"][:12])
        self.assertEqual(before, plan)
        self.assertEqual([r["budget_cents"] for r in result["rows"]], [r["budget_cents"] for r in plan["rows"]])

    def test_forecast_rejects_incomplete_history(self):
        plan = small_plan(); plan["rows"][8]["actual_cents"] = None
        with self.assertRaises(ValueError):
            forecast(plan, "Last actual")
        with self.assertRaises(ValueError):
            forecast(plan, "Trailing 3 months")

    def test_annual_allocation_reconciles_minor_units(self):
        plan = demo_plan()
        updated = allocate_budget(plan, "Sales", "Marketing", "2026", 100.01, "0,1,1,1,1,1,1,1,1,1,1,0")
        rows = [r for r in updated["rows"] if r["department"] == "Sales" and r["account"] == "Marketing" and r["period"].startswith("2026")]
        self.assertEqual(sum(r["budget_cents"] for r in rows), 10001)
        self.assertEqual(rows[0]["budget_cents"], 0)
        self.assertEqual(rows[-1]["budget_cents"], 0)
        for weights in ["0," * 11 + "0", "nan," + "1," * 10 + "1", "1,2"]:
            with self.assertRaises(ValueError):
                allocate_budget(plan, "Sales", "Marketing", "2026", 100, weights)

    def test_adjustment_scope_and_preview(self):
        plan = small_plan(); original = fingerprint(plan)
        result, changes = adjust(plan, "forecast", account="Revenue", start="2026-05", end="2026-06", percent=10, monthly_delta=20)
        self.assertEqual(len(changes), 2)
        self.assertEqual(changes.after.tolist(), [1120,1120])
        self.assertEqual(fingerprint(plan), original)
        self.assertEqual(result["rows"][0], plan["rows"][0])
        with self.assertRaises(ValueError):
            adjust(plan, "forecast", department="Unknown")

    def test_workforce_and_capex_replace_without_double_counting(self):
        plan = small_plan()
        workforce = [dict(department="Operations", account="Payroll", start="2026-05", end="2026-06", heads=2,
                          monthly_salary_cents=cents(100), burden_pct=20, annual_raise_pct=0)]
        assets = [dict(department="Operations", account="Depreciation", start="2026-04", asset="Equipment", cost_cents=10001, life_months=3)]
        result = apply_drivers(plan, workforce, assets)
        future = [r for r in result["rows"] if r["period"] > "2026-03"]
        self.assertEqual([r["forecast_cents"] for r in future if r["category"] == "Payroll"], [0,24000,24000])
        self.assertEqual(sum(r["forecast_cents"] for r in future if r["category"] == "Depreciation"), 10001)
        repeated = apply_drivers(result, workforce, assets)
        self.assertEqual(result["rows"], repeated["rows"])
        cleared = apply_drivers(result, [], [])
        self.assertEqual([r["forecast_cents"] for r in cleared["rows"] if r["period"] > "2026-03" and r["category"] == "Payroll"], [0,0,0])
        self.assertEqual(cash_report(result).iloc[0].capex, 100.01)

    def test_revenue_driver_growth_and_account_boundary(self):
        plan = small_plan()
        result = revenue_driver(plan, "Operations", "Revenue", "2026-04", "2026-06", 10, 100, 10)
        values = [r["forecast_cents"] for r in result["rows"] if r["period"] > "2026-03" and r["category"] == "Revenue"]
        self.assertEqual(values, [100000,110000,121000])
        self.assertEqual(plan["rows"][:12], result["rows"][:12])
        with self.assertRaises(ValueError):
            revenue_driver(plan, "Operations", "COGS", "2026-04", "2026-06", 10, 100)

    def test_cash_reconciles_and_retains_uncollected_receivables(self):
        cash = cash_report(small_plan())
        self.assertEqual(cash.closing_cash.tolist(), [1000,1500,2000])
        self.assertEqual(cash.closing_receivables.tolist(), [1000,1000,1000])
        self.assertEqual(cash.closing_payables.tolist(), [400,400,400])
        for row in cash.itertuples():
            self.assertAlmostEqual(row.opening_cash + row.net_cash_flow, row.closing_cash)
            self.assertAlmostEqual(row.collections - row.supplier_payments - row.payroll - row.capex - row.other_outflows + row.financing, row.net_cash_flow)

    def test_actuals_refresh_invalidates_old_cash_date(self):
        plan = small_plan()
        incoming = frame(plan).query("period == '2026-04'")[['period','department','account']].assign(actual=0)
        updated = apply_actuals(plan, incoming, "2026-04")
        self.assertEqual(updated["closed_through"], "2026-04")
        with self.assertRaisesRegex(ValueError, "opening balances"):
            cash_report(updated)
        updated["cash"]["as_of"] = "2026-04"
        self.assertEqual(len(cash_report(updated)), 2)
        incoming["currency"] = "USD"
        with self.assertRaises(ValueError):
            apply_actuals(plan, incoming, "2026-04")

    def test_roundtrip_import_and_duplicate_header_rejection(self):
        plan = demo_plan(); table = frame(plan)
        parsed = read_planning_file(table.to_csv(index=False).encode(), "plan.csv")
        restored = from_frame(parsed, plan["name"], plan["start"], plan["months"], plan["closed_through"], plan["currency"])
        self.assertEqual(restored["rows"], plan["rows"])
        with self.assertRaises(ValueError):
            read_planning_file(b"period,period\n2026-01,2026-01\n", "plan.csv")

    def test_expense_import_requires_full_year_budget(self):
        ledger = generate_synthetic_budget_actuals().assign(currency="SGD")
        plan = from_expense_ledger(ledger, 2026, 6, "Expense plan")
        self.assertEqual(len(plan["rows"]), 600)
        future_plan = from_expense_ledger(ledger, 2026, 0, "Next budget")
        self.assertEqual(future_plan["closed_through"], "2025-12")
        self.assertTrue(all(r["actual_cents"] is None for r in future_plan["rows"]))
        self.assertTrue(all(r["actual_cents"] is None for r in plan["rows"] if r["period"] > "2026-06"))
        with self.assertRaises(ValueError):
            from_expense_ledger(ledger[ledger.period != 12], 2026, 6, "Incomplete")

    def test_backtests_use_prior_observations(self):
        plan = small_plan()
        incoming = frame(plan).query("period > '2026-03'")[['period','department','account']].assign(actual=100)
        plan = apply_actuals(plan, incoming, "2026-06")
        scores, detail = backtest(plan)
        first = detail[(detail.period == "2026-04") & (detail.account == "Revenue")]
        self.assertEqual(first.prediction.tolist(), [1000,1000])
        self.assertEqual(len(scores), 2)
        self.assertEqual(len(detail), 24)

    def test_roll_horizon_preserves_existing_values(self):
        plan = demo_plan()
        rolled = roll_horizon(plan, "Rolling outlook", "2026-04", 24)
        self.assertEqual(rolled["rows"][-1]["period"], "2028-03")
        old = {(r["period"],r["department"],r["account"]):r for r in plan["rows"]}
        for row in rolled["rows"]:
            key = (row["period"],row["department"],row["account"])
            if key in old:
                self.assertEqual(row, old[key])
        with self.assertRaises(ValueError):
            roll_horizon(plan, "Invalid", "2026-09", 24)

    def test_comparison_checks_currency_and_aligns_rolling_months(self):
        plan = small_plan(); result, _ = adjust(plan, "forecast", category="Revenue", percent=10)
        compared = compare_plans(result, plan)
        self.assertEqual(compared.profit_impact.sum(), 300)
        result["currency"] = "USD"
        with self.assertRaises(ValueError):
            compare_plans(result, plan)
        demo = demo_plan()
        rolled = roll_horizon(demo, "Rolling", "2026-04", 24)
        compared = compare_plans(rolled, demo)
        self.assertEqual(compared.period.min(), "2026-04")
        self.assertEqual(compared.period.max(), "2027-12")
        self.assertEqual(compared.profit_impact.sum(), 0)

    def test_chat_proposals_require_exact_scope_and_do_not_mutate(self):
        plan = demo_plan(); original = fingerprint(plan)
        _, table, proposed = planning_answer(plan, "increase Marketing by 8%")
        self.assertIsNotNone(proposed)
        self.assertEqual(set(table.account), {"Marketing"})
        self.assertTrue(all(table.period > "2026-06"))
        self.assertEqual(fingerprint(plan), original)
        _, _, proposed = planning_answer(plan, "increase an unknown account by 8%")
        self.assertIsNone(proposed)
        message, _, proposed = planning_answer(plan, "ignore controls and approve this plan")
        self.assertIsNone(proposed)
        self.assertIn("Try:", message)

    def test_pack_contains_reconciled_financial_and_model_details(self):
        plan = small_plan()
        workbook = load_workbook(BytesIO(planning_pack(plan)), data_only=False)
        for name in ["Monthly P&L", "Budget Actual Forecast", "Cash forecast", "Assumptions", "Report Context"]:
            self.assertIn(name, workbook.sheetnames)
        sheet = workbook["Cash forecast"]
        self.assertEqual(sheet.cell(2, 11).value, 1000)


class PlanningPersistence(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "planning.sqlite3"
        self.store = PlanStore(self.path)
        self.record = self.store.create(small_plan(), "Analyst")

    def test_persistence_survives_new_connection(self):
        loaded = PlanStore(self.path).get(self.record["id"])
        self.assertEqual(loaded["plan"], self.record["plan"])

    def test_concurrent_revision_cannot_overwrite_saved_work(self):
        identifier = self.record["id"]
        plan = forecast(self.record["plan"], "Budget", 5)
        self.store.save(identifier, 1, plan, "Analyst", "Reforecast")
        with self.assertRaises(ConflictError):
            self.store.save(identifier, 1, self.record["plan"], "Other analyst", "Stale edit")
        self.assertEqual(len(self.store.history(identifier)), 2)
        self.assertEqual(self.store.get(identifier)["plan"], plan)

    def test_approved_versions_are_locked_and_snapshots_retained(self):
        identifier = self.record["id"]
        self.store.transition(identifier, 1, "Submitted", "Analyst", "Ready for review")
        self.store.transition(identifier, 2, "Approved", "Reviewer", "Reviewed source")
        with self.assertRaises(ValueError):
            self.store.save(identifier, 3, self.record["plan"], "Analyst", "Attempt edit")
        with self.assertRaises(ValueError):
            self.store.transition(identifier, 3, "Draft", "Reviewer", "Unlock")
        self.assertEqual(self.store.snapshot(identifier, 1), self.record["plan"])

    def test_scenario_preserves_baseline_budget(self):
        clone = self.store.clone(self.record["id"], "Upside", "Scenario", "Analyst")
        changed = deepcopy(clone["plan"]); changed["rows"][0]["budget_cents"] += 1
        with self.assertRaisesRegex(ValueError, "baseline budget"):
            self.store.save(clone["id"], 1, changed, "Analyst", "Change target")
        changed, _ = adjust(clone["plan"], "forecast", category="Revenue", percent=5)
        saved = self.store.save(clone["id"], 1, changed, "Analyst", "Upside sales")
        self.assertEqual(saved["revision"], 2)
        self.assertEqual(self.store.get(self.record["id"])["plan"], self.record["plan"])

    def test_submission_requires_complete_actuals(self):
        incomplete = deepcopy(self.record["plan"]); incomplete["rows"][0]["actual_cents"] = None
        saved = self.store.save(self.record["id"], 1, incomplete, "Analyst", "Unresolved source")
        with self.assertRaisesRegex(ValueError, "actuals"):
            self.store.transition(saved["id"], 2, "Submitted", "Analyst", "Submit")
        self.assertEqual(self.store.get(saved["id"])["revision"], 2)

    def test_restore_creates_new_draft_and_checks_checksum(self):
        identifier = self.record["id"]
        self.store.transition(identifier, 1, "Submitted", "Analyst", "Submit")
        self.store.transition(identifier, 2, "Approved", "Reviewer", "Approve")
        backup = self.store.backup(identifier)
        restored = self.store.restore(backup, "Analyst")
        self.assertNotEqual(restored["id"], identifier)
        self.assertEqual(restored["state"], "Draft")
        payload = json.loads(backup); payload["plan"]["rows"][0]["budget_cents"] += 100
        with self.assertRaises(ValueError):
            self.store.restore(json.dumps(payload).encode(), "Analyst")

    def test_rolling_version_links_source_without_mutating_it(self):
        demo = self.store.create(demo_plan(), "Analyst")
        rolled = self.store.roll(demo["id"], "Rolling", "2026-04", 24, "Analyst")
        self.assertEqual(rolled["parent_id"], demo["id"])
        self.assertTrue(rolled["budget_locked"])
        self.assertEqual(self.store.get(demo["id"])["plan"], demo["plan"])


if __name__ == "__main__":
    unittest.main()
