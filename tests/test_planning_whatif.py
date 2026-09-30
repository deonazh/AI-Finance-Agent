"""Budget/forecast what-if scope, accounting effects, cash timing, and exports."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from epm import demo_plan, fingerprint
from planning_chatbot import PlanningChatbot


PROMPT = 'What if I increase Marketing budget by 8% for next month?'


class PlanningWhatIfTests(unittest.TestCase):
    def setUp(self):
        self.plan = demo_plan()
        self.today = date(2026, 9, 30)
        self.bot = PlanningChatbot(self.plan, today=self.today)

    def answer(self, question=PROMPT, plan=None, **kwargs):
        return PlanningChatbot(plan or self.plan, today=self.today).answer(question, **kwargs)

    def impact(self, result, metric, report='Impact summary'):
        return result.reports[report].set_index('metric').loc[metric, 'change']

    def test_exact_question_changes_one_budget_month_only(self):
        before = deepcopy(self.plan)
        result = self.answer()
        self.assertIsNotNone(result.proposal)
        self.assertEqual(self.plan, before)
        self.assertEqual(len(result.table), 1)
        self.assertEqual(result.table.iloc[0].period, '2026-10')
        self.assertEqual(result.table.iloc[0].before, 15600)
        self.assertEqual(result.table.iloc[0].after, 16848)
        self.assertIn('October 2026', result.message)
        self.assertIn('2026-09-30', result.message)
        for old, new in zip(before['rows'], result.proposal['rows']):
            if old['period'] == '2026-10' and old['account'] == 'Marketing':
                self.assertEqual(new['budget_cents'] - old['budget_cents'], 124800)
            else:
                self.assertEqual(old['budget_cents'], new['budget_cents'])
            self.assertEqual(old['actual_cents'], new['actual_cents'])
            self.assertEqual(old['forecast_cents'], new['forecast_cents'])
        self.assertIn('Nothing has been saved', result.message)

    def test_budget_and_illustrative_spend_effects_are_separate(self):
        result = self.answer()
        self.assertEqual(self.impact(result, 'Company budget operating profit'), -1248)
        self.assertEqual(self.impact(result, 'Company forecast operating profit'), 0)
        self.assertEqual(self.impact(result, 'Company favorable profit variance'), 1248)
        self.assertEqual(self.impact(result, 'Company closing cash at 2027-12'), 0)
        alt = 'Budget and spend impact'
        self.assertEqual(self.impact(result, 'Company forecast operating profit', alt), -1248)
        self.assertEqual(self.impact(result, 'Company favorable profit variance', alt), 0)
        self.assertEqual(self.impact(result, 'Company closing cash at 2027-12', alt), -1248)
        self.assertIn('No extra sales', result.message)

    def test_spending_alternative_adds_increment_to_existing_forecast(self):
        for row in self.plan['rows']:
            if row['account'] == 'Marketing' and row['period'] == '2026-10':
                row['forecast_cents'] = 2000000
        result = self.answer()
        alt = result.reports['Budget and spend impact'].set_index('metric')
        self.assertEqual(alt.loc['Forecast: selected Opex lines', 'after'], 21248)
        self.assertEqual(alt.loc['Forecast: selected Opex lines', 'change'], 1248)

    def test_cash_follows_payment_lag_and_reconciles_payables(self):
        for lag, first in [(0, '2026-10'), (1, '2026-11'), (3, '2027-01')]:
            with self.subTest(lag=lag):
                self.plan['cash']['payment_lag'] = lag
                result = self.answer()
                cash = result.reports['Budget and spend cash'].set_index('period')
                affected = cash[cash.net_cash_flow_change.abs() > .005]
                self.assertEqual(affected.index.tolist(), [first])
                self.assertEqual(affected.iloc[0].net_cash_flow_change, -1248)
                self.assertEqual(cash.loc['2026-10', 'closing_payables_change'], 1248 if lag else 0)
                self.assertEqual(cash.loc[first, 'closing_payables_change'], 0)

    def test_horizon_boundary_reports_unpaid_cost_not_fabricated_cash(self):
        result = self.answer('increase Marketing forecast by 8% for December 2027')
        cash = result.reports['Monthly cash impact'].iloc[-1]
        self.assertEqual(cash.closing_cash_change, 0)
        self.assertEqual(cash.closing_payables_change, 1248)
        self.assertEqual(self.impact(result, 'Company forecast operating profit'), -1248)
        self.assertIn('No cash-flow effect falls inside the horizon', result.message)

    def test_stale_cash_still_allows_profit_preview_without_cash_claim(self):
        self.plan['cash']['as_of'] = '2026-05'
        result = self.answer()
        self.assertIsNotNone(result.proposal)
        self.assertTrue(result.reports['Monthly cash impact'].empty)
        self.assertNotIn('Closing cash at', result.message)
        self.assertTrue(any('Cash impact unavailable' in w for w in result.warnings))

    def test_revenue_and_non_cash_accounts_have_correct_signs(self):
        revenue = self.answer('increase Service contracts budget by 8% for October 2026')
        self.assertEqual(self.impact(revenue, 'Company budget operating profit'), 14976)
        self.assertEqual(self.impact(revenue, 'Company forecast operating profit', 'Budget and spend impact'), 14976)
        self.assertEqual(self.impact(revenue, 'Company closing cash at 2027-12', 'Budget and spend impact'), 14976)
        depreciation = self.answer('increase Depreciation forecast by 8% for October 2026')
        self.assertLess(self.impact(depreciation, 'Company forecast operating profit'), 0)
        self.assertEqual(self.impact(depreciation, 'Company closing cash at 2027-12'), 0)

    def test_payroll_scope_can_select_one_department(self):
        result = self.answer('increase Sales / Payroll forecast by 8% for October 2026')
        self.assertEqual(result.table.department.tolist(), ['Sales'])
        cash = result.reports['Monthly cash impact']
        self.assertEqual(cash[cash.net_cash_flow_change.ne(0)].period.tolist(), ['2026-10'])
        second = self.answer('increase Payroll in Sales forecast by 8% for October 2026')
        pd.testing.assert_frame_equal(result.table, second.table)

    def test_expense_only_plan_labels_profit_limit(self):
        self.plan['rows'] = [r for r in self.plan['rows'] if r['category'] != 'Revenue']
        result = self.answer()
        self.assertTrue(any('no revenue' in w for w in result.warnings))

    def test_relative_month_is_calendar_based_and_handles_year_boundary(self):
        next_month = self.answer()
        first_open = self.answer('increase Marketing budget by 8% for first forecast month')
        self.assertEqual(next_month.context['change']['start'], '2026-10')
        self.assertEqual(first_open.context['change']['start'], '2026-07')
        self.today = date(2026, 12, 31)
        self.assertEqual(self.answer().context['change']['start'], '2027-01')

    def test_explicit_period_forms_cover_exact_requested_months(self):
        for phrase, expected in [('October 2026', 1), ('2026-10', 1), ('Q4 2026', 3), ('2027', 12), ('next year', 12), ('2026-10 to 2027-01', 4)]:
            with self.subTest(phrase=phrase):
                result = self.answer('increase Marketing budget by 8% for ' + phrase)
                self.assertIsNotNone(result.proposal, result.message)
                self.assertEqual(len(result.table), expected)

    def test_natural_wrappers_percent_and_field_position(self):
        queries = ['What would happen if we increased the Marketing budget by 8 percent for next month?',
                   'Can you increase budget for Marketing by 8% for next month?',
                   'increase Marketing budget for next month by 8%',
                   'please raise Marketing budget by 8 per cent for next month']
        expected = self.answer().table
        for query in queries:
            with self.subTest(query=query):
                result = self.answer(query)
                self.assertIsNotNone(result.proposal, result.message)
                pd.testing.assert_frame_equal(result.table, expected)
        reduced = self.answer('What would happen if we reduced Marketing forecast by 5 percent in November 2026?')
        self.assertEqual(reduced.table.iloc[0].after, 14820)

    def test_legacy_forecast_command_labels_default_open_horizon(self):
        result = self.answer('increase Marketing by 8%')
        self.assertEqual(len(result.table), 18)
        self.assertEqual(result.context['change']['field'], 'forecast')
        self.assertIn('all open months', result.message)
        for old, new in zip(self.plan['rows'], result.proposal['rows']):
            self.assertEqual(old['budget_cents'], new['budget_cents'])

    def test_invalid_scope_period_or_extra_instructions_never_broaden(self):
        queries = [
            'increase Marketing budget by 8% for October',
            'increase Marketing budget by 8% for June 2026',
            'increase Marketing budget by 8% for 2026',
            'increase Marketing budget by 8% for 2028',
            'increase Marketing budget by 8% for 2026-13',
            'increase Marketing budget by 8% for 2026-10 to 2026-09',
            'increase Marketing budget by 8% for 2027-12 to 2028-02',
            'increase Marketing budget by 8% for next month and Sales by 3%',
            'increase Atlantis budget by 8% for next month',
            'increase Marketing and Payroll budget by 8% for next month',
            'increase Marketing budget and forecast by 8% for next month',
            'increase Marketing budget by 1001% for next month',
            'decrease Marketing budget by 101% for next month',
            'increase Marketing budget by -8% for next month',
            'increase Marketing budget by 8% for next month and save it',
            'what if I boost Marketing budget by 8% for next month',
        ]
        for query in queries:
            with self.subTest(query=query):
                result = self.answer(query)
                self.assertIsNone(result.proposal)
                self.assertIsNone(result.table)
                self.assertEqual(result.context, {})

    def test_outdated_demo_does_not_reinterpret_next_month(self):
        self.today = date(2028, 1, 1)
        result = self.answer()
        self.assertIsNone(result.proposal)
        self.assertIn('2028-02', result.message)
        self.assertIn('outside this plan', result.message)

    def test_followup_reuses_rate_field_scope_and_allows_inherited_year(self):
        first = self.answer()
        second = self.answer('what about November?', context=first.context)
        self.assertEqual(second.context['change']['start'], '2026-11')
        self.assertEqual(second.context['change']['field'], 'budget')
        self.assertEqual(second.context['change']['account'], 'Marketing')
        self.assertEqual(second.context['change']['percent'], 8)
        self.assertEqual(len(second.table), 1)

    def test_export_reconciles_preview_and_labels_illustration(self):
        first = self.answer()
        output = self.answer('export that to Excel', context=first.context)
        pd.testing.assert_frame_equal(first.table, output.table)
        book = load_workbook(BytesIO(output.export_bytes), data_only=True)
        self.assertIn('Budget and spend impact', book.sheetnames)
        values = list(book['Proposed line changes'].values)
        self.assertEqual(dict(zip(values[0], values[1]))['change'], 1248)
        summary = {r[0]: r[-1] for r in list(book['Impact summary'].values)[1:]}
        self.assertEqual(summary['Company forecast operating profit'], 0)
        alt = {r[0]: r[-1] for r in list(book['Budget and spend impact'].values)[1:]}
        self.assertEqual(alt['Company forecast operating profit'], -1248)
        metadata = dict(book['Report Context'].values)
        self.assertEqual(metadata['Plan fingerprint'], fingerprint(self.plan))
        self.assertEqual(metadata['Status'], 'Preview only; no changes saved')
        self.assertIsNone(self.answer('export that for 2027', context=first.context).export_bytes)

    def test_changed_revision_discards_previous_whatif_scope(self):
        first = self.answer()
        self.plan['rows'][0]['budget_cents'] += 100
        result = self.answer('export that to Excel', context=first.context)
        self.assertIsNone(result.export_bytes)
        self.assertIsNone(result.proposal)

    def test_retained_request_is_revalidated(self):
        initial = self.answer().context
        for key, value in [('percent', float('inf')), ('percent', True), ('percent', 1001), ('start', '2026-01'), ('account', 'Atlantis'), ('field', 'actual')]:
            with self.subTest(key=key, value=value):
                context = deepcopy(initial); context['change'][key] = value
                self.assertIsNone(self.answer('export that to Excel', context=context).proposal)

    def test_reserved_all_name_cannot_expand_adjustment_scope(self):
        for row in self.plan['rows']:
            if row['account'] == 'Marketing':
                row['account'] = 'All'
        result = self.answer('increase account All budget by 8% for next month')
        self.assertIsNone(result.proposal)

    def test_rounded_budget_increment_reconciles_to_forecast_cents(self):
        for row in self.plan['rows']:
            if row['account'] == 'Marketing' and row['period'] == '2026-10':
                row['budget_cents'] = 12345
        result = self.answer()
        self.assertAlmostEqual(result.table.iloc[0]['change'], 9.88)
        self.assertAlmostEqual(self.impact(result, 'Company forecast operating profit', 'Budget and spend impact'), -9.88)

    @patch('planning_chatbot._get_chat_model')
    def test_whatifs_never_send_financial_questions_to_provider(self, factory):
        for question in [PROMPT, 'what if I boost Marketing budget by 8%', 'increase Atlantis budget by 8%']:
            self.answer(question, provider='openai')
        factory.assert_not_called()


if __name__ == '__main__':
    unittest.main()
