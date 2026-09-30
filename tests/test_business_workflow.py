"""Business data integrity, review state, and reporting export contracts."""
from io import BytesIO
import json
import unittest
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

from agents import _build_deterministic_analysis, executive_storyteller_agent, run_variance_analysis
from chatbot import FinanceDataChatbot
from data_engine import load_budget_actuals_file, calculate_variances, filter_significant_variances
from exports import csv_bytes, workbook_bytes
from guardrails import validate_response_against_dataframe
from review import ledger_id, review_register, save_edits, review_json, restore_review, month_end_pack


def ledger():
    return pd.DataFrame([
        dict(fiscal_year=2026, period=1, cost_center="Finance", gl_account="Cloud Hosting", budget=1000.0, actual=1500.0, currency="SGD"),
        dict(fiscal_year=2026, period=2, cost_center="Finance", gl_account="Cloud Hosting", budget=1000.0, actual=900.0, currency="SGD"),
        dict(fiscal_year=2026, period=2, cost_center="HR", gl_account="Recruitment", budget=0.0, actual=500.0, currency="SGD"),
    ])


def upload(frame):
    file = BytesIO(frame.to_csv(index=False).encode()); file.name = "ledger.csv"
    return load_budget_actuals_file(file)


class UploadValidationTests(unittest.TestCase):
    def test_valid_upload_recalculates_stale_derived_fields(self):
        data = upload(ledger().assign(variance=999, quarter="Q4", period_label="stale"))
        self.assertEqual(data.variance.tolist(), [500, -100, 500])
        self.assertEqual(set(data.quarter), {"Q1"})
        self.assertEqual(data.period_label.iloc[0], "FY2026-P01")

    def test_invalid_numbers_are_never_coerced_to_zero(self):
        for value in ["unknown", "", float("inf"), float("nan"), "$1,000"]:
            with self.subTest(value=value):
                data = ledger(); data["budget"] = data["budget"].astype(object); data.loc[0,"budget"] = value
                with self.assertRaisesRegex(ValueError,"budget must contain numeric"):
                    upload(data)

    def test_periods_and_years_must_be_valid_integers(self):
        for col, value in [("period",0),("period",13),("period",1.5),("fiscal_year",0),("fiscal_year",2026.5)]:
            with self.subTest(col=col,value=value):
                data = ledger(); data[col] = data[col].astype(float); data.loc[0,col] = value
                with self.assertRaisesRegex(ValueError,"whole number"):
                    upload(data)

    def test_duplicates_and_blank_dimensions_are_rejected(self):
        with self.assertRaisesRegex(ValueError,"Duplicate monthly"):
            upload(pd.concat([ledger(),ledger().iloc[:1]]))
        for col in ["cost_center","gl_account"]:
            with self.subTest(column=col):
                with self.assertRaisesRegex(ValueError,"cannot be blank"):
                    upload(ledger().assign(**{col:" "}))

    def test_alias_collision_and_empty_file_are_rejected(self):
        with self.assertRaisesRegex(ValueError,"Duplicate column"):
            upload(ledger().assign(budget_amount=0))
        with self.assertRaisesRegex(ValueError,"no ledger rows"):
            upload(ledger().iloc[:0])

    def test_duplicate_raw_headers_are_rejected_for_csv_and_xlsx(self):
        data=ledger(); data.insert(len(data.columns), "second_budget", 1)
        data.columns=["budget" if name=="second_budget" else name for name in data.columns]
        for kind in ("csv", "xlsx"):
            file=BytesIO();file.name="ledger."+kind
            if kind=="csv":file.write(data.to_csv(index=False).encode())
            else:data.to_excel(file,index=False)
            file.seek(0)
            with self.subTest(kind=kind),self.assertRaisesRegex(ValueError,"Duplicate column"):
                load_budget_actuals_file(file)

    def test_finite_inputs_cannot_overflow_reporting(self):
        import numpy as np
        with np.errstate(over='ignore',invalid='ignore'),self.assertRaises(ValueError):
            upload(ledger().assign(budget=1e308,actual=1e308))

    def test_mixed_and_missing_currencies_are_rejected(self):
        data = ledger(); data.loc[0,"currency"] = "USD"
        with self.assertRaisesRegex(ValueError,"Mixed currencies"):
            upload(data)
        data.loc[0,"currency"] = ""
        with self.assertRaisesRegex(ValueError,"currency must"):
            upload(data)

    def test_credits_are_retained_and_equal_spend_is_on_budget(self):
        data=upload(ledger().assign(budget=-100,actual=-100))
        self.assertEqual(set(data.variance_direction),{"On budget"})
        self.assertEqual(data.budget.sum(),-300)

    def test_unbudgeted_activity_is_always_flagged(self):
        data=upload(ledger())
        flagged=filter_significant_variances(data,50,50000)
        self.assertEqual(len(flagged),1)
        self.assertEqual(flagged.iloc[0].flag_reason,"Unbudgeted activity")
        self.assertTrue(pd.isna(flagged.iloc[0].variance_pct))

    def test_thresholds_cannot_be_negative_or_nonfinite(self):
        for pct,amount in [(-1,100),(10,-1),(float('nan'),10)]:
            with self.assertRaises(ValueError): filter_significant_variances(ledger(),pct,amount)


class BusinessAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.raw=upload(ledger())
        self.bot=FinanceDataChatbot(self.raw,filter_significant_variances(self.raw),self.raw)

    def test_zero_budget_workflow_produces_valid_brief_and_chat(self):
        raw=self.raw[self.raw.budget==0]
        brief=run_variance_analysis(raw,use_llm=False)
        self.assertIn("N/A (zero budget)",brief.executive_markdown)
        self.assertIsNone(brief.analyst_json['net_variance_pct'])
        json.dumps(brief.analyst_json,allow_nan=False)
        bot=FinanceDataChatbot(raw,filter_significant_variances(raw),raw)
        self.assertIn("S$500",bot.answer("biggest variance",use_llm=False).message)
        self.assertIn("N/A",bot.answer("biggest variance",use_llm=False).message)

    def test_material_row_count_is_separate_from_reviewed_count(self):
        result=_build_deterministic_analysis(self.raw,10,50000)
        self.assertEqual(result.reviewed_row_count,3)
        self.assertEqual(result.flagged_row_count,2)

    def test_account_name_does_not_prove_root_cause(self):
        result=_build_deterministic_analysis(self.raw,10,50000)
        self.assertTrue(all(d.driver_type=='unknown' for d in result.top_cost_drivers))
        self.assertEqual(self.bot._infer_driver_type("Finance","Cloud Hosting"),'unknown')

    def test_stale_variance_columns_are_recalculated_in_chat(self):
        raw=self.raw.assign(variance=9_999_999)
        bot=FinanceDataChatbot(raw,raw,raw)
        answer=bot.answer('total variance',use_llm=False)
        self.assertIn('S$900',answer.message)
        self.assertNotIn('9,999,999',answer.message)

    def test_comparison_does_not_treat_missing_month_as_zero(self):
        response=self.bot.answer('compare January vs December costs',use_llm=False)
        self.assertEqual(response.agent_used,'Guardrail_Agent')
        self.assertIn('Both periods',response.message)

    def test_wrong_year_is_not_silently_ignored(self):
        response=self.bot.answer('total variance FY2025',use_llm=False)
        self.assertEqual(response.agent_used,'Guardrail_Agent')
        self.assertIn('2026',response.message)

    def test_large_department_counts_are_supported(self):
        raw=pd.concat([self.raw.assign(cost_center=f'Department {i}') for i in range(15)],ignore_index=True)
        result=run_variance_analysis(raw,use_llm=False)
        self.assertEqual(len(result.analyst_json['cost_center_summaries']),15)

    def test_storyteller_rejects_unverified_numbers(self):
        analysis=_build_deterministic_analysis(self.raw,10,50000).model_dump()
        model=Mock()
        model.with_structured_output.return_value.invoke.return_value={'title':'Brief','markdown':'Actuals were $987,654,321.'}
        with patch('agents._get_chat_model',return_value=model):
            result=executive_storyteller_agent({'analyst_json':analysis,'use_llm':True})
        self.assertNotIn('987,654,321',result['executive_markdown'])
        self.assertTrue(result['errors'])

    def test_guardrail_checks_small_currency_amounts_and_empty_answers(self):
        self.assertFalse(validate_response_against_dataframe('Cost was $99.',{'actual':500}))
        self.assertFalse(validate_response_against_dataframe('',{'actual':500}))
        self.assertTrue(validate_response_against_dataframe('1. Cost was $500.',{'actual':500}))

    def test_excel_profile_reconciles_to_filtered_detail(self):
        response=self.bot.answer('export Finance to Excel',use_llm=False)
        book=load_workbook(BytesIO(response.export_bytes),data_only=True)
        values=list(book['Profile'].values); profile=dict(zip(values[0],values[1]))
        self.assertEqual(profile['rows'],2)
        self.assertEqual(profile['total_actual'],2400)
        self.assertEqual(profile['currency'],'SGD')
        self.assertEqual(book['Detail'].freeze_panes,'A2')


class ReviewWorkflowTests(unittest.TestCase):
    def setUp(self): self.raw=upload(ledger())

    def test_review_notes_roundtrip_and_reject_different_ledger(self):
        register=review_register(self.raw)
        register.loc[0,['owner','status','explanation']]=['Finance lead','Resolved','Confirmed against invoice']
        notes=save_edits(register,{},self.raw)
        backup=review_json(self.raw,notes)
        self.assertEqual(restore_review(backup,self.raw),notes)
        self.assertEqual(ledger_id(self.raw.iloc[::-1]),ledger_id(self.raw))
        with self.assertRaisesRegex(ValueError,'different ledger'):
            restore_review(backup,self.raw.assign(actual=0))

    def test_resolved_notes_need_owner_and_evidence(self):
        register=review_register(self.raw);register.loc[0,'status']='Resolved'
        with self.assertRaisesRegex(ValueError,'owner and an explanation'):
            save_edits(register,{},self.raw)

    def test_filtered_edits_preserve_other_review_notes(self):
        register=review_register(self.raw);register['owner']='A'
        notes=save_edits(register,{},self.raw)
        edited=register.iloc[:1].copy();edited['owner']='B'
        updated=save_edits(edited,notes,self.raw)
        self.assertEqual(len(updated),len(notes))
        self.assertEqual(updated[register.iloc[-1].review_id]['owner'],'A')

    def test_invalid_dates_and_status_are_rejected(self):
        for col,value in [('due_date','yesterday'),('status','Approved')]:
            register=review_register(self.raw);register.loc[0,col]=value
            with self.assertRaises(ValueError):save_edits(register,{},self.raw)

    def test_report_pack_contains_reconciling_totals_and_context(self):
        content=month_end_pack(self.raw,review_register(self.raw),{'Source':'synthetic.csv'})
        book=load_workbook(BytesIO(content),data_only=True)
        self.assertEqual(book.sheetnames,['Department Summary','Ledger Detail','Review Register','Methodology','Report Context'])
        context=dict(list(book['Report Context'].values)[1:])
        self.assertEqual(context['Total actual'],2900)
        self.assertEqual(context['Variance'],900)
        self.assertEqual(context['Currency'],'SGD')

    def test_formula_text_remains_literal_in_excel_and_csv(self):
        data=pd.DataFrame({'text':['=1+1',' +SUM(A1)','@SUM(A1)','-text'],'actual':[-1.5,2,3,4]})
        book=load_workbook(BytesIO(workbook_bytes({'Data':data})),data_only=False)
        self.assertEqual(book['Data']['A2'].value,'=1+1')
        self.assertEqual(book['Data']['A2'].data_type,'s')
        csv=csv_bytes(data).decode('utf-8-sig')
        self.assertIn("'=1+1",csv)
        self.assertIn('-1.5',csv)


class ReviewEmptyFieldTests(unittest.TestCase):
    def test_null_fields_do_not_resolve_an_unassigned_item(self):
        from review import validate_notes
        with self.assertRaisesRegex(ValueError,'requires an owner'):
            validate_notes({'item':{'owner':None,'status':'Resolved','explanation':None}}, {'item'})

    def test_cleared_optional_fields_are_empty_text(self):
        from review import validate_notes
        result=validate_notes({'item':{'owner':None,'due_date':float('nan'),'explanation':None}}, {'item'})
        self.assertEqual(result['item']['owner'],'')
        self.assertEqual(result['item']['due_date'],'')
        self.assertEqual(result['item']['status'],'Open')
