"""Regression tests for business-style chatbot prompts."""

from __future__ import annotations

import json
from datetime import datetime
import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from chatbot import FinanceDataChatbot
from chatbot import _fetch_exchange_rate
from data_engine import filter_significant_variances, generate_synthetic_budget_actuals
from prompts import SYSTEM_PROMPT, TOOL_DESCRIPTIONS


class MockHTTPResponse:
    def __init__(self, payload: bytes) -> None:
        self.payload = payload

    def __enter__(self) -> "MockHTTPResponse":
        return self

    def __exit__(self, *args: object) -> None:
        return None

    def read(self) -> bytes:
        return self.payload


class FinanceDataChatbotPromptTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        raw_df = generate_synthetic_budget_actuals()
        flagged_df = filter_significant_variances(raw_df)
        cls.bot = FinanceDataChatbot(raw_df=raw_df, flagged_df=flagged_df, current_view_df=flagged_df)

    def setUp(self):
        clock = patch("chatbot.datetime")
        clock.start().now.return_value = datetime(2026, 9, 12)
        self.addCleanup(clock.stop)
        provider = patch("chatbot._get_chat_model", return_value=None)
        provider.start()
        self.addCleanup(provider.stop)

    def test_system_prompt_contains_strict_enterprise_guardrails(self) -> None:
        self.assertIn("DETERMINISTIC MATH RULE", SYSTEM_PROMPT)
        self.assertIn("ZERO DATA HALLUCINATION", SYSTEM_PROMPT)
        self.assertIn("query_financial_dataset", SYSTEM_PROMPT)
        self.assertIn("calculate_variance_metrics", SYSTEM_PROMPT)

    def test_langchain_tool_aliases_are_available(self) -> None:
        tools = self.bot._build_tools()  # pylint: disable=protected-access
        names = {tool.name for tool in tools}

        self.assertIn("query_financial_dataset", names)
        self.assertIn("calculate_variance_metrics", names)
        self.assertIn("convert_currency", names)
        self.assertIn("export_excel_report", names)
        self.assertIn(TOOL_DESCRIPTIONS["calculate_variance_metrics"], {tool.description for tool in tools})

    def test_calculate_variance_metrics_tool_returns_python_calculated_values(self) -> None:
        payload = json.loads(
            self.bot._tool_calculate_variance_metrics(  # pylint: disable=protected-access
                source="raw",
                cost_center="Marketing",
                gl_account="Marketing Campaigns",
                period=10,
                group_by="total",
            )
        )

        self.assertEqual(payload["calculation_basis"], "Python calculated: Variance_USD = Actual - Budget; Variance_Pct = Variance_USD / Budget * 100.")
        self.assertAlmostEqual(
            payload["variance_usd"],
            payload["total_actual"] - payload["total_budget"],
            places=2,
        )

    def test_biggest_variance_in_december_is_not_glossary(self) -> None:
        response = self.bot.answer("what is the biggest variance in december")

        self.assertIn("biggest variance", response.message.lower())
        self.assertIn("period = 12", response.message)
        self.assertNotIn("Actual minus budget. Positive is overspend", response.message)

    def test_biggest_variance_in_q4_uses_quarter_filter(self) -> None:
        response = self.bot.answer("what is the biggest variance in Q4")

        self.assertIn("biggest variance", response.message.lower())
        self.assertIn("quarter = Q4", response.message)

    def test_top_variances_returns_ranked_table(self) -> None:
        response = self.bot.answer("top 5 variances in Q4")

        self.assertIn("top variance", response.message.lower())
        self.assertIn("quarter = Q4", response.message)
        self.assertIn("| budget | actual | variance |", response.message)

    def test_largest_cost_center_routes_to_cost_center_grain(self) -> None:
        response = self.bot.answer("which cost center has the biggest variance")

        self.assertIn("biggest variance by `cost_center`", response.message)
        self.assertNotIn("Actual minus budget. Positive is overspend", response.message)

    def test_compare_defaults_to_raw_ledger(self) -> None:
        response = self.bot.answer("compare budget vs actual for Marketing by account")

        self.assertIn("budget vs actual", response.message.lower())
        self.assertIn("`raw`", response.message)
        self.assertIn("cost center = Marketing", response.message)

    def test_total_variance_defaults_to_raw_ledger(self) -> None:
        response = self.bot.answer("total variance in Q4")

        self.assertIn("For `raw`, total variance", response.message)
        self.assertIn("quarter = Q4", response.message)

    def test_spend_question_defaults_to_raw_ledger(self) -> None:
        response = self.bot.answer("how much did Marketing spend in December")

        self.assertIn("Total actual spend for `raw`", response.message)
        self.assertIn("period = 12", response.message)
        self.assertIn("cost center = Marketing", response.message)

    def test_column_definition_still_works_when_asked_directly(self) -> None:
        response = self.bot.answer("what is variance_pct?")

        self.assertIn("Variance divided by budget", response.message)

    def test_upload_schema_question_does_not_trigger_excel_export(self) -> None:
        response = self.bot.answer("what data must i have in my excel upload file?")

        self.assertIsNone(response.export_bytes)
        self.assertIn("Required columns", response.message)
        self.assertIn("`fiscal_year`", response.message)
        self.assertIn("`.csv` or `.xlsx`", response.message)
        self.assertEqual(response.agent_used, "Upload_Schema_Agent")

    def test_upload_schema_question_clears_stale_export(self) -> None:
        self.bot._last_export = {  # pylint: disable=protected-access
            "bytes": b"old workbook",
            "filename": "old.xlsx",
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }

        response = self.bot.answer("what columns do I need in my xlsx upload file?")

        self.assertIsNone(response.export_bytes)
        self.assertIsNone(response.export_filename)
        self.assertIsNone(self.bot._last_export)  # pylint: disable=protected-access

    def test_invalid_period_is_rejected(self) -> None:
        response = self.bot.answer("what is the biggest variance in period 13")

        self.assertIn("outside the supported range", response.message)

    def test_out_of_scope_question_does_not_guess(self) -> None:
        response = self.bot.answer("who approved the invoice for Marketing Payroll")

        self.assertIn("cannot answer that reliably", response.message)
        self.assertIn("does not contain invoice approvers", response.message)

    def test_delivery_request_explains_missing_integration(self) -> None:
        response = self.bot.answer("please email this report to CFO")

        self.assertIn("cannot answer that reliably", response.message)
        self.assertIn("delivery credentials", response.message)

    def test_personal_identity_question_is_declined_not_profiled(self) -> None:
        response = self.bot.answer("am i gay")

        self.assertEqual(response.agent_used, "Guardrail_Agent")
        self.assertIn("Enterprise Financial Analytics Assistant", response.message)
        self.assertNotIn("current_view", response.message)

    def test_general_off_topic_question_is_declined(self) -> None:
        response = self.bot.answer("can you write a python script to scrape movie ratings?")

        self.assertEqual(response.agent_used, "Guardrail_Agent")
        self.assertIn("budget analysis", response.message)
        self.assertNotIn("current_view", response.message)

    def test_greeting_does_not_show_current_view_profile(self) -> None:
        response = self.bot.answer("hello")

        self.assertEqual(response.agent_used, "Greeting_Agent")
        self.assertIn("Finance Data Copilot", response.message)
        self.assertIn("how do I use this app?", response.message)
        self.assertNotIn("current_view", response.message)

    def test_how_to_use_app_returns_end_to_end_guide(self) -> None:
        response = self.bot.answer("how do I use this app?")

        self.assertEqual(response.agent_used, "Guide_Agent")
        self.assertIn("end-to-end guide", response.message)
        self.assertIn("Load Data", response.message)
        self.assertIn("Run AI Analysis", response.message)
        self.assertIn("Use The Chatbot", response.message)
        self.assertNotIn("current_view", response.message.splitlines()[0])

    def test_how_ai_analysis_works_returns_trust_explanation(self) -> None:
        response = self.bot.answer("how does the AI analysis work?")

        self.assertEqual(response.agent_used, "Guide_Agent")
        self.assertIn("Python calculates all numbers first", response.message)
        self.assertIn("The app does not let the LLM perform raw arithmetic", response.message)
        self.assertIn("Variance is calculated in Python", response.message)

    def test_can_i_trust_analysis_returns_guardrail_explanation(self) -> None:
        response = self.bot.answer("can I trust the analysis?")

        self.assertEqual(response.agent_used, "Guide_Agent")
        self.assertIn("Calculation Controls", response.message)
        self.assertIn("Guardrails reject unsupported periods", response.message)

    def test_ambiguous_quarter_performance_question_routes_to_finance_summary(self) -> None:
        response = self.bot.answer("how are we doing this quarter?")

        self.assertEqual(response.agent_used, "Executive_Briefing_Agent")
        self.assertIn("Management finance summary", response.message)
        self.assertIn("**Executive Summary**", response.message)
        self.assertIn("quarter = Q", response.message)

    def test_boss_presentation_last_month_summary_is_executive_snapshot(self) -> None:
        response = self.bot.answer("I have a presentation to my boss in one hour, summarise last month finances for me")

        self.assertEqual(response.agent_used, "Executive_Briefing_Agent")
        self.assertIn("Management finance summary", response.message)
        self.assertIn("August (P08)", response.message)
        self.assertIn("**Total Budget**", response.message)
        self.assertIn("**Top Variance Drivers**", response.message)
        self.assertIn("**Rows Reviewed**: 50", response.message)
        self.assertNotIn("cost center = Finance", response.message)
        self.assertNotIn("top variance groups from `current_view`", response.message.lower())

    def test_last_month_finance_summary_uses_raw_august_scope(self) -> None:
        response = self.bot.answer("summarise last month finances")

        self.assertEqual(response.agent_used, "Executive_Briefing_Agent")
        self.assertIn("from `raw`", response.message)
        self.assertIn("period = 8", response.message)
        self.assertIn("August (P08)", response.message)
        self.assertIn("**Rows Reviewed**: 50", response.message)
        self.assertNotIn("cost center = Finance", response.message)

    def test_best_cost_saving_last_month_routes_to_favorable_driver(self) -> None:
        response = self.bot.answer("which was our best cost saving last month?")

        self.assertEqual(response.agent_used, "Variance_Investigation_Agent")
        self.assertIn("largest cost saving driver", response.message)
        self.assertIn("in `raw`", response.message)
        self.assertIn("period = 8", response.message)
        self.assertIn("direction = Favorable", response.message)
        self.assertNotIn("current_view", response.message)

    def test_where_are_we_overspending_last_month_lists_unfavorable_drivers(self) -> None:
        response = self.bot.answer("where are we overspending last month?")

        self.assertEqual(response.agent_used, "Data_Analyst_Agent")
        self.assertIn("top overspends", response.message)
        self.assertIn("period = 8", response.message)
        self.assertIn("direction = Unfavorable", response.message)
        self.assertIn("| cost_center | gl_account |", response.message)
        self.assertNotIn("| period_label | budget | actual |", response.message)
        self.assertNotIn("current_view", response.message)

    def test_which_cost_center_saved_most_last_month_groups_by_cost_center(self) -> None:
        response = self.bot.answer("which cost center saved the most last month?")

        self.assertEqual(response.agent_used, "Variance_Investigation_Agent")
        self.assertIn("largest cost saving by `cost_center`", response.message)
        self.assertIn("period = 8", response.message)
        self.assertIn("direction = Favorable", response.message)
        self.assertNotIn("current_view", response.message)

    def test_how_much_did_we_save_last_month_totals_favorable_variance(self) -> None:
        response = self.bot.answer("how much did we save last month?")

        self.assertEqual(response.agent_used, "Data_Analyst_Agent")
        self.assertIn("total cost savings", response.message.lower())
        self.assertIn("period = 8", response.message)
        self.assertIn("direction = Favorable", response.message)
        self.assertNotIn("not confident", response.message)

    def test_how_much_did_we_overspend_last_month_totals_unfavorable_variance(self) -> None:
        response = self.bot.answer("how much did we overspend last month?")

        self.assertEqual(response.agent_used, "Data_Analyst_Agent")
        self.assertIn("total overspend", response.message.lower())
        self.assertIn("period = 8", response.message)
        self.assertIn("direction = Unfavorable", response.message)

    def test_are_we_over_budget_last_month_uses_net_scope(self) -> None:
        response = self.bot.answer("are we over budget last month?")

        self.assertEqual(response.agent_used, "Executive_Briefing_Agent")
        self.assertIn("Management finance summary", response.message)
        self.assertIn("period = 8", response.message)
        self.assertNotIn("direction = Unfavorable", response.message)

    def test_anomalies_last_month_routes_to_driver_table(self) -> None:
        response = self.bot.answer("any anomalies last month?")

        self.assertEqual(response.agent_used, "Data_Analyst_Agent")
        self.assertIn("top anomalies", response.message)
        self.assertIn("period = 8", response.message)
        self.assertIn("| cost_center | gl_account |", response.message)

    def test_month_to_month_business_language_routes_to_comparison(self) -> None:
        response = self.bot.answer("what changed between october and november?")

        self.assertEqual(response.agent_used, "Data_Analyst_Agent")
        self.assertIn("cost movement from October (P10) to November (P11)", response.message)
        self.assertIn("periods = October (P10), November (P11)", response.message)
        self.assertNotIn("current_view", response.message)

    def test_why_did_marketing_overspend_routes_to_investigation(self) -> None:
        response = self.bot.answer("why did Marketing overspend in November?")

        self.assertEqual(response.agent_used, "Variance_Investigation_Agent")
        self.assertIn("largest overspend driver", response.message)
        self.assertIn("cost center = Marketing", response.message)
        self.assertIn("period = 11", response.message)
        self.assertIn("direction = Unfavorable", response.message)

    def test_management_actions_for_overspends_get_action_view(self) -> None:
        response = self.bot.answer("what actions should we take on overspends last month?")

        self.assertEqual(response.agent_used, "Executive_Action_Agent")
        self.assertIn("management action view", response.message)
        self.assertIn("Priority Drivers & Actions", response.message)
        self.assertIn("period = 8", response.message)
        self.assertIn("direction = Unfavorable", response.message)

    def test_show_departments_below_budget_groups_by_cost_center(self) -> None:
        response = self.bot.answer("show departments below budget in Q2")

        self.assertEqual(response.agent_used, "Data_Analyst_Agent")
        self.assertIn("top cost savings", response.message)
        self.assertIn("quarter = Q2", response.message)
        self.assertIn("direction = Favorable", response.message)
        self.assertIn("| cost_center | budget | actual | variance |", response.message)
        self.assertNotIn("| cost_center | gl_account |", response.message)

    def test_give_me_accounts_above_budget_groups_by_gl_account(self) -> None:
        response = self.bot.answer("give me accounts above budget in December")

        self.assertEqual(response.agent_used, "Data_Analyst_Agent")
        self.assertIn("top overspends", response.message)
        self.assertIn("period = 12", response.message)
        self.assertIn("direction = Unfavorable", response.message)
        self.assertIn("| gl_account | budget | actual | variance |", response.message)
        self.assertNotIn("| cost_center | gl_account |", response.message)

    def test_biggest_variance_investigation_excel_has_expected_sheets(self) -> None:
        response = self.bot.answer("give me excel report of the biggest variance investigation in december")

        self.assertIsNotNone(response.export_bytes)
        self.assertTrue(response.export_filename.endswith(".xlsx"))
        workbook = load_workbook(BytesIO(response.export_bytes), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Investigation", "Full Year Detail", "Flagged Rows", "Monthly Trend", "Top Drivers"],
        )

    def test_period_comparison_excel_uses_both_months(self) -> None:
        response = self.bot.answer("excel report of difference in marketing cost between october and november")

        self.assertIsNotNone(response.export_bytes)
        self.assertIn("October (P10) to November (P11)", response.message)
        self.assertIn("cost center = Marketing", response.message)
        self.assertIn("periods = October (P10), November (P11)", response.message)
        workbook = load_workbook(BytesIO(response.export_bytes), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Comparison Summary", "Driver Breakdown", "Period Detail", "Monthly Trend"],
        )

    def test_followup_inherits_previous_compare_context(self) -> None:
        history = [
            {"role": "user", "content": "compare budget vs actual for Marketing by account"},
            {"role": "assistant", "content": "Here is the budget vs actual from `raw`. Filters applied: cost center = Marketing."},
        ]

        response = self.bot.answer("what about December?", chat_history=history)

        self.assertIn("budget vs actual", response.message.lower())
        self.assertIn("`raw`", response.message)
        self.assertIn("cost center = Marketing", response.message)
        self.assertIn("period = 12", response.message)

    def test_followup_export_inherits_biggest_variance_context(self) -> None:
        first = self.bot.answer("what is the biggest variance in December?")
        history = [
            {"role": "user", "content": "what is the biggest variance in December?"},
            {"role": "assistant", "content": first.message},
        ]

        response = self.bot.answer("export that to Excel", chat_history=history)

        self.assertIsNotNone(response.export_bytes)
        self.assertIn("focused Excel investigation workbook", response.message)
        self.assertTrue(response.export_filename.endswith(".xlsx"))

    def test_exchange_rate_uses_online_tool_result(self) -> None:
        with patch(
            "chatbot._fetch_exchange_rate",
            return_value={
                "base_currency": "USD",
                "quote_currency": "SGD",
                "amount": 100.0,
                "rate": 1.35,
                "converted_amount": 135.0,
                "date": "2026-09-09",
                "source": "Frankfurter public exchange-rate API",
            },
        ):
            response = self.bot.answer("convert 100 USD to SGD")

        self.assertIn("100.00 USD = 135.00 SGD", response.message)
        self.assertIn("Frankfurter", response.message)

    def test_exchange_rate_parser_supports_frankfurter_rate_shape(self) -> None:
        payload = b'{"date":"2026-09-09","base":"USD","quote":"SGD","rate":1.2657}'
        with patch("chatbot.urlopen", return_value=MockHTTPResponse(payload)):
            result = _fetch_exchange_rate("USD", "SGD", 100)

        self.assertEqual(result["base_currency"], "USD")
        self.assertEqual(result["quote_currency"], "SGD")
        self.assertEqual(result["converted_amount"], 126.57)


if __name__ == "__main__":
    unittest.main()
