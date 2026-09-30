"""Scope and tool-completion regressions without live provider calls."""
import unittest
from io import BytesIO
from unittest.mock import Mock, patch

from langchain_core.messages import AIMessage
from openpyxl import load_workbook

from chatbot import FinanceDataChatbot
from data_engine import generate_synthetic_budget_actuals, filter_significant_variances


class CopilotWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.raw = generate_synthetic_budget_actuals()
        self.flagged = filter_significant_variances(self.raw)
        self.view = self.raw[(self.raw.cost_center == "Finance") & (self.raw.period == 1)]
        self.bot = FinanceDataChatbot(self.raw, self.flagged, self.view, default_source="current_view")

    def test_dashboard_scope_applies_to_totals(self):
        response = self.bot.answer("what is the total actual spend?", use_llm=False)
        self.assertIn("`current_view`", response.message)
        self.assertIn(f"${self.view.actual.sum():,.0f}", response.message)

    def test_explicit_full_ledger_overrides_dashboard_scope(self):
        response = self.bot.answer("total variance for the full ledger", use_llm=False)
        self.assertIn("`raw`", response.message)
        self.assertIn(f"${(self.raw.actual.sum()-self.raw.budget.sum()):,.0f}", response.message)

    def test_scope_applies_to_executive_summary(self):
        response = self.bot.answer("summarise finances", use_llm=False)
        self.assertIn(f"${self.view.actual.sum():,.0f}", response.message)
        self.assertNotIn(f"${self.raw.actual.sum():,.0f}", response.message)

    def test_export_uses_exact_dashboard_rows(self):
        response = self.bot.answer("export to Excel", use_llm=False)
        workbook = load_workbook(BytesIO(response.export_bytes), data_only=True)
        self.assertEqual(workbook["Detail"].max_row, len(self.view) + 1)

    def test_empty_scope_never_falls_back_to_full_ledger(self):
        bot = FinanceDataChatbot(self.raw, self.flagged, self.view.iloc[:0], default_source="current_view")
        response = bot.answer("top 5 unfavorable variances", use_llm=False)
        self.assertEqual(response.agent_used, "Guardrail_Agent")
        self.assertIsNone(response.export_bytes)

    def model_with_answer(self, content):
        model = Mock()
        model.bind_tools.return_value = model
        model.invoke.side_effect = [
            AIMessage(content="", tool_calls=[{"name": "calculate_variance_metrics", "id": "metrics_1", "args":
                {"source": "current_view", "group_by": "total"}}]),
            AIMessage(content=content),
        ]
        return model

    def test_valid_tool_backed_answer_finishes_without_more_calls(self):
        content = f"Actual spend is ${self.view.actual.sum():,.2f}."
        model = self.model_with_answer(content)
        with patch("chatbot._get_chat_model", return_value=model):
            response = self.bot.answer("what is the total actual spend?", use_llm=True)
        self.assertTrue(response.used_llm)
        self.assertEqual(response.message, content)
        self.assertEqual(model.invoke.call_count, 2)

    def test_unverified_final_numbers_still_trigger_local_fallback(self):
        model = self.model_with_answer("Actual spend is $987,654,321.00.")
        with patch("chatbot._get_chat_model", return_value=model):
            response = self.bot.answer("what is the total actual spend?", use_llm=True)
        self.assertFalse(response.used_llm)
        self.assertNotIn("987,654,321", response.message)
        self.assertIn(f"${self.view.actual.sum():,.0f}", response.message)

    def test_answer_without_data_tool_uses_local_fallback(self):
        model = Mock()
        model.bind_tools.return_value = model
        model.invoke.return_value = AIMessage(content="Actual spend is $987,654,321.00.")
        with patch("chatbot._get_chat_model", return_value=model):
            response = self.bot.answer("what is the total actual spend?", use_llm=True)
        self.assertFalse(response.used_llm)
        self.assertNotIn("987,654,321", response.message)

    def test_full_conversation_does_not_inherit_accounts_from_answer_table(self):
        bot = FinanceDataChatbot(self.raw, self.flagged, self.raw, default_source="raw")
        history = []
        for question in ["compare budget vs actual for Marketing by account", "what about December?", "export that to Excel"]:
            response = bot.answer(question, use_llm=False, chat_history=history)
            history += [{"role": "user", "content": question}, {"role": "assistant", "content": response.message}]
        workbook = load_workbook(BytesIO(response.export_bytes), data_only=True)
        rows = list(workbook["Detail"].values)
        header = rows[0]
        self.assertEqual(len(rows)-1, 10)
        self.assertEqual({row[header.index("period")] for row in rows[1:]}, {12})
        self.assertEqual({row[header.index("cost_center")] for row in rows[1:]}, {"Marketing"})

    def test_followup_month_replaces_previous_quarter(self):
        bot = FinanceDataChatbot(self.raw, self.flagged, self.raw, default_source="raw")
        history = [{"role": "user", "content": "compare budget vs actual for Marketing in Q4 by account"}]
        response = bot.answer("what about February?", use_llm=False, chat_history=history)
        self.assertIn("period = 2", response.message)
        self.assertNotIn("quarter = Q4", response.message)
        self.assertIn("budget vs actual", response.message)

    def test_followup_respects_previous_explicit_full_ledger(self):
        history = [{"role": "user", "content": "compare budget vs actual in full ledger for Marketing by account"}]
        response = self.bot.answer("what about December?", use_llm=False, chat_history=history)
        self.assertIn("`raw`", response.message)
        self.assertIn("period = 12", response.message)

    def test_new_topic_stops_filters_from_older_topic(self):
        bot = FinanceDataChatbot(self.raw, self.flagged, self.raw, default_source="raw")
        history = [{"role": "user", "content": "compare budget vs actual for Marketing Campaigns in Q4"},
                   {"role": "user", "content": "compare budget vs actual for Finance by account"}]
        response = bot.answer("what about December?", use_llm=False, chat_history=history)
        self.assertIn("cost center = Finance", response.message)
        self.assertNotIn("gl account = Marketing Campaigns", response.message)
        self.assertNotIn("quarter = Q4", response.message)

    def test_export_after_month_change_does_not_revive_old_quarter(self):
        bot = FinanceDataChatbot(self.raw, self.flagged, self.raw, default_source="raw")
        history = [{"role": "user", "content": "compare budget vs actual for Marketing in Q4 by account"},
                   {"role": "user", "content": "what about February?"}]
        response = bot.answer("export that to Excel", use_llm=False, chat_history=history)
        self.assertIsNotNone(response.export_bytes)
        workbook = load_workbook(BytesIO(response.export_bytes), data_only=True)
        self.assertEqual(workbook["Detail"].max_row, 11)
        self.assertIn("period = 2", response.message)
        self.assertNotIn("quarter = Q4", response.message)
