"""Financial and scope regressions for the plan-aware assistant."""
from copy import deepcopy
from io import BytesIO
import unittest
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import load_workbook
from epm import demo_plan, fingerprint, monthly_report, cash_report
from planning_chatbot import PlanningChatbot, ReportIntent


class PlanningChatTests(unittest.TestCase):
    def setUp(self):
        self.plan = demo_plan()
        self.bot = PlanningChatbot(self.plan)

    def test_budget_department_and_year_reconcile_to_inputs(self):
        result = self.bot.answer("show budget for Sales in 2027")
        self.assertEqual(len(result.table), 12)
        expected = sum(r['budget_cents'] for r in self.plan['rows'] if r['department']=='Sales' and r['period'].startswith('2027') and r['category']=='Revenue')/100
        self.assertAlmostEqual(result.table.budget_revenue.sum(), expected)
        self.assertEqual(result.context['department'], 'Sales')

    def test_followup_and_export_preserve_validated_scope(self):
        first = self.bot.answer('show budget for Sales in 2026')
        second = self.bot.answer('what about 2027?', context=first.context)
        output = self.bot.answer('export that to Excel', context=second.context)
        self.assertEqual(output.context['department'], 'Sales')
        pd.testing.assert_frame_equal(second.table, output.table)
        book = load_workbook(BytesIO(output.export_bytes), data_only=True)
        self.assertEqual(book['Planning answer'].max_row, 13)
        metadata = dict(book['Report Context'].values)
        self.assertEqual(metadata['Start'], '2027-01')
        self.assertEqual(metadata['Plan fingerprint'], fingerprint(self.plan))

    def test_fresh_question_resets_filters_but_navigation_can_keep_context(self):
        first = self.bot.answer('show budget for Sales in 2027')
        second = self.bot.answer('summarise plan', context=first.context)
        self.assertNotIn('department', second.context)
        self.assertEqual(len(second.table), 24)

    def test_context_does_not_cross_plan_revisions(self):
        first = self.bot.answer('show budget for Sales in 2027')
        self.plan['rows'][0]['budget_cents'] += 100
        response = PlanningChatbot(self.plan).answer('export that to Excel', context=first.context)
        self.assertIsNone(response.export_bytes)
        self.assertIn('report first', response.message)

    def test_unknown_or_ambiguous_scope_never_broadens(self):
        for query in ['show budget for Atlantis', 'show budget for Sales and Corporate', 'show budget for July', 'show budget for Q3', 'show budget for 2030', 'show budget for next year', 'show forecast from 2026-09 to 2026-07']:
            with self.subTest(query=query):
                result = self.bot.answer(query)
                self.assertIsNone(result.table)
                self.assertEqual(result.context, {})

    def test_forecast_excludes_closed_actuals_and_honors_quarter(self):
        result = self.bot.answer('show forecast for Q3 2026')
        self.assertEqual(list(result.table.period), ['2026-07','2026-08','2026-09'])
        self.assertIn('forecast_operating_profit', result.table)
        self.assertIsNone(self.bot.answer('show forecast for 2026-13').table)

    def test_missing_actuals_remain_unknown_in_summary(self):
        self.plan['rows'][0]['actual_cents'] = None
        result = PlanningChatbot(self.plan).answer('summarise plan')
        self.assertTrue(pd.isna(result.table.iloc[0].outlook_operating_profit))
        self.assertIn('unknown', result.message)

    def test_expense_scope_is_not_described_as_company_profit(self):
        result = self.bot.answer('summarise Marketing')
        self.assertIn('negative costs', result.message)
        self.assertTrue((result.table.outlook_opex > 0).all())

    def test_natural_cost_question_reports_cost_total(self):
        result = self.bot.answer("how much is Payroll in 2027?")
        expected = sum(r['forecast_cents'] for r in self.plan['rows'] if r['category']=='Payroll' and r['period'].startswith('2027'))/100
        self.assertAlmostEqual(result.table.outlook_payroll.sum(), expected)
        self.assertIn(f"SGD {expected:,.2f}", result.message)
        self.assertIn('not total company profitability', result.message)

    def test_cash_filters_after_full_horizon_calculation(self):
        result = self.bot.answer('show cash for 2027')
        expected = cash_report(self.plan).query("period >= '2027-01'").reset_index(drop=True)
        pd.testing.assert_frame_equal(result.table, expected)
        self.assertIsNone(self.bot.answer('show cash for Sales').table)

    def test_stale_cash_balances_are_not_reported(self):
        self.plan['cash']['as_of'] = '2026-05'
        result = PlanningChatbot(self.plan).answer('show cash')
        self.assertIsNone(result.table)
        self.assertIn('earlier close', result.message)

    def test_schedule_amounts_are_reporting_currency(self):
        result = self.bot.answer('show workforce')
        self.assertNotIn('monthly_salary_cents', result.table)
        self.assertEqual(result.table.iloc[0].monthly_salary, self.plan['workforce'][0]['monthly_salary_cents']/100)
        assets = self.bot.answer('show capex for Corporate')
        self.assertEqual(set(assets.table.department), {'Corporate'})

    def test_accuracy_evaluation_honors_periods_without_future_leakage(self):
        result = self.bot.answer('forecast accuracy for 2026-04 to 2026-05')
        self.assertEqual(set(result.table.period), {'2026-04','2026-05'})
        changed = deepcopy(self.plan)
        for row in changed['rows']:
            if row['period'] > '2026-05':
                row['actual_cents'] = 99999999 if row['actual_cents'] is not None else None
        repeated = PlanningChatbot(changed).answer('forecast accuracy for 2026-04 to 2026-05')
        pd.testing.assert_frame_equal(result.table, repeated.table)

    def test_variances_report_only_unfavorable_known_rows(self):
        result = self.bot.answer('top variances for 2027')
        self.assertLessEqual(len(result.table), 10)
        self.assertTrue((result.table.favorable_impact < 0).all())
        self.assertTrue(result.table.period.str.startswith('2027').all())

    def test_proposal_never_mutates_input_or_closed_periods(self):
        original = deepcopy(self.plan)
        result = self.bot.answer('increase Marketing by 8%')
        self.assertIsNotNone(result.proposal)
        self.assertEqual(self.plan, original)
        self.assertEqual(self.bot.plan, original)
        for before, after in zip(original['rows'],result.proposal['rows']):
            self.assertEqual(before['budget_cents'],after['budget_cents'])
            self.assertEqual(before['actual_cents'],after['actual_cents'])
            if before['period'] <= original['closed_through']:
                self.assertEqual(before,after)

    @patch('planning_chatbot._get_chat_model')
    def test_model_cannot_route_writes_or_unknown_explicit_scopes(self, model):
        for q in ['approve this plan', 'delete all data', 'increase Atlantis by 9%', 'show budget for Atlantis']:
            result = self.bot.answer(q, provider='openai')
            self.assertIsNone(result.proposal)
        model.assert_not_called()

    @patch('planning_chatbot._get_chat_model')
    def test_optional_provider_only_classifies_and_python_calculates(self, factory):
        factory.return_value.with_structured_output.return_value.invoke.return_value = ReportIntent(report='summary')
        result = self.bot.answer('give me a financial snapshot', provider='openai')
        self.assertTrue(result.used_llm)
        expected = monthly_report(self.plan).outlook_operating_profit
        pd.testing.assert_series_equal(result.table.outlook_operating_profit, expected)
        messages = factory.return_value.with_structured_output.return_value.invoke.call_args.args[0]
        self.assertEqual(messages[1][1], 'give me a financial snapshot')
        self.assertNotIn(self.plan['name'], str(messages))

    @patch('planning_chatbot._get_chat_model')
    def test_invalid_provider_output_fails_closed_and_redacts_errors(self, factory):
        for value in [{'report':'summary','amount':999999}, {'report':'delete'}]:
            factory.return_value.with_structured_output.return_value.invoke.return_value = value
            result = self.bot.answer('financial snapshot please', provider='openai')
            self.assertFalse(result.used_llm)
            self.assertIsNone(result.table)
            self.assertTrue(result.warnings)
        factory.side_effect = RuntimeError('secret-token')
        result = self.bot.answer('financial snapshot please', provider='openai')
        self.assertNotIn('secret-token', str(result))

    @patch('planning_chatbot._get_chat_model', return_value=None)
    def test_missing_provider_and_local_prompts_work(self, factory):
        result = self.bot.answer('financial snapshot please', provider='openai')
        self.assertTrue(result.warnings)
        factory.reset_mock()
        self.assertIsNotNone(self.bot.answer('show budget',provider='openai').table)
        factory.assert_not_called()

    def test_empty_and_oversized_questions_are_bounded(self):
        for question in ['', 'x'*4001, None]:
            self.assertIsNone(self.bot.answer(question).table)


if __name__ == '__main__':
    unittest.main()
