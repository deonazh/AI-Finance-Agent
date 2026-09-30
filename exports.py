"""In-memory reports with readable formatting and literal spreadsheet text."""
from io import BytesIO
from datetime import datetime, timezone

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def csv_bytes(df):
    """Neutralise spreadsheet formulas in text without changing numeric amounts."""
    def literal(value):
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return value
    return df.map(literal).to_csv(index=False).encode("utf-8-sig")


def format_workbook(workbook):
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="15253C")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # openpyxl otherwise treats text beginning '=' as executable formula.
                if cell.data_type == "f":
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str))
                if row[0].row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F1F5F9")
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);–'
        for col in sheet.columns:
            lengths = [len(str(c.value or "")) for c in list(col)[:100]]
            sheet.column_dimensions[get_column_letter(col[0].column)].width = min(48, max(14, max(lengths, default=0) + 2))


def workbook_bytes(sheets, metadata=None):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
        if metadata is not None:
            values = {"Generated (UTC)": datetime.now(timezone.utc).isoformat(timespec="seconds"), **metadata}
            pd.DataFrame(list(values.items()), columns=["Field", "Value"]).to_excel(writer, sheet_name="Report Context", index=False)
        format_workbook(writer.book)
    return buffer.getvalue()
